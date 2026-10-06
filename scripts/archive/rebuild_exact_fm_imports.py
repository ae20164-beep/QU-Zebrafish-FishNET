import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

labels_dir = r'c:\Users\ae20164\OneDrive - Qatar University (1)\Zebrafish shared folder\FishNET Data\Labels'
export_dir = r'c:\Users\ae20164\OneDrive - Qatar University (1)\Zebrafish shared folder\FishNET Data\Labels\FishNet Exported Data'
fishnet_dir = r'c:\Users\ae20164\OneDrive - Qatar University (1)\Zebrafish shared folder\FishNET Data'


# 1. Load FishNET.tab master data
fn_tab_path = os.path.join(export_dir, 'FishNET.tab')
if not os.path.exists(fn_tab_path):
    fn_tab_path = os.path.join(labels_dir, 'FishNET.tab')

with open(fn_tab_path, 'r', encoding='utf-8-sig', errors='ignore') as f:
    fn_rows = list(csv.reader(f, delimiter='\t'))

fn_headers = fn_rows[0]
fn_tanks = [dict(zip(fn_headers, r)) for r in fn_rows[1:] if r and r[0].strip()]
fn_tanks_by_id = {t['TUID'].strip().upper(): t for t in fn_tanks}

def clean_val(v):
    if v is None: return ""
    return str(v).replace('\x0b', ' / ').replace('\x0c', ' ').replace('\x00', '').strip()

def get_genotype_desc(notes):
    n = notes.lower().strip()
    if 'fli' in n: return 'Tg (fli1a:eGFP) Sidra [Fli]'
    elif 'gata' in n:
        if 'ab' in n or 'sidra' in n or 'x' in n:
            return 'Tg (gata1:dsRed) Sidra [Gata]; Wt (AB)'
        return 'Tg (gata1:dsRed) Sidra [Gata]'
    elif 'casper' in n or 'cas' in n: return 'Mu (mitfaw2/w2; mpv17a9/a9) [Casper]'
    elif 'ab' in n or 'wild' in n or 'wt' in n: return 'Wt (AB)'
    return 'Wt (AB)'

def get_line_nick(notes):
    n = notes.lower().strip()
    if 'casper' in n or 'cas' in n: return 'Casper'
    if 'fli' in n: return 'Fli'
    if 'gata' in n: return 'Gata'
    if 'ab' in n or 'wt' in n or 'wild' in n: return 'AB'
    return 'AB'

# =========================================================================
# 1. CROSSES TABLE - EXACT FILEMAKER FIELD NAMES
# =========================================================================
# Exact FileMaker fields:
# CUID, Paternal ID, Maternal ID, Set up by, Date of Birth, Date Started, Status, Protocol, Lab, Number of Tanks, Number of Pairs set up, Desired Offspring, 0HPF SR, 24HPF SR, # survived 0HPF, # survived 24HPF, # of dead 0HPF, # of dead 24HPF, Notes

crosses_field_names = [
    'CUID',
    'Paternal ID',
    'Maternal ID',
    'Set up by',
    'Date of Birth',
    'Date Started',
    'Status',
    'Protocol',
    'Lab',
    'Number of Tanks',
    'Number of Pairs set up',
    'Desired Offspring',
    '0HPF SR',
    '24HPF SR',
    '# survived 0HPF',
    '# survived 24HPF',
    '# of dead 0HPF',
    '# of dead 24HPF',
    'Notes'
]

# Read existing crosses.tab.orig.bak
with open(os.path.join(export_dir, 'crosses.tab.orig.bak'), 'r', encoding='utf-8-sig', errors='ignore') as f:
    orig_crosses_raw = list(csv.reader(f, delimiter='\t'))

crosses_records = {}
for r in orig_crosses_raw:
    if len(r) > 7 and r[7].strip():
        cuid = r[7].strip().upper()
        # orig columns: 0:dead0, 1:dead24, 2:tot, 3:viable, 4:sr0, 5:sr24, 6:tech, 7:cuid, 8:d_log, 9:sp_date, 10:setup_date, 11:mat, 12:trans, 13:'', 14:pairs, 15:pat, 16:iacuc, 17:pi, 18:stat, 19:lab, 20:mat_geno, 21:mat_notes, 22:pat_geno, 23:pat_notes
        crosses_records[cuid] = [
            cuid,                               # CUID
            r[15] if len(r)>15 else '',         # Paternal ID
            r[11] if len(r)>11 else '',         # Maternal ID
            r[6] if len(r)>6 else 'Enas',       # Set up by
            r[8] if len(r)>8 else '',           # Date of Birth
            r[10] if len(r)>10 else '',         # Date Started
            r[18] if len(r)>18 else 'Active',   # Status
            r[16] if len(r)>16 else 'QU-IACUC 006/2023-AMM5', # Protocol
            r[19] if len(r)>19 else 'Zebrafish facility',     # Lab
            r[14] if len(r)>14 else '1',        # Number of Tanks
            r[14] if len(r)>14 else '1',        # Number of Pairs set up
            r[2] if len(r)>2 else '14',         # Desired Offspring
            r[4] if len(r)>4 else '100.00%',    # 0HPF SR
            r[5] if len(r)>5 else '100.00%',    # 24HPF SR
            r[3] if len(r)>3 else '14',         # # survived 0HPF
            r[3] if len(r)>3 else '14',         # # survived 24HPF
            r[0] if len(r)>0 else '0',          # # of dead 0HPF
            r[1] if len(r)>1 else '0',          # # of dead 24HPF
            (r[21] + ' x ' + r[23]) if len(r)>23 else '' # Notes
        ]

# Add new crosses from FishNET.tab
cross_to_tanks = {}
for t in fn_tanks:
    c = t['Derivative cross'].strip().upper()
    if c:
        if c not in cross_to_tanks: cross_to_tanks[c] = []
        cross_to_tanks[c].append(t)

for cuid, t_list in cross_to_tanks.items():
    if cuid not in crosses_records:
        first_t = t_list[0]
        pat = first_t['PATERNAL'].strip()
        mat = first_t['MATERNAL'].strip()
        tot = str(sum(int(t['TOTAL']) for t in t_list if t['TOTAL'].isdigit()))
        dob = first_t['DOB'].strip()
        notes = first_t['NOTES'].strip()
        
        crosses_records[cuid] = [
            cuid,                               # CUID
            pat if pat else mat,                # Paternal ID
            mat if mat else pat,                # Maternal ID
            'Enas',                             # Set up by
            dob,                                # Date of Birth
            dob,                                # Date Started
            'Active',                           # Status
            'QU-IACUC 006/2023-AMM5',           # Protocol
            'Zebrafish facility',               # Lab
            '1',                                # Number of Tanks
            '1',                                # Number of Pairs set up
            tot if tot != '0' else '14',        # Desired Offspring
            '100.00%',                          # 0HPF SR
            '100.00%',                          # 24HPF SR
            tot if tot != '0' else '14',        # # survived 0HPF
            tot if tot != '0' else '14',        # # survived 24HPF
            '0',                                # # of dead 0HPF
            '0',                                # # of dead 24HPF
            notes                               # Notes
        ]

# Sort crosses
import re
def cuid_sort_key(k):
    m = re.search(r'\d+', k)
    return int(m.group(0)) if m else 9999

sorted_crosses_rows = [crosses_records[k] for k in sorted(crosses_records.keys(), key=cuid_sort_key)]

# Save Crosses Excel & CSV with EXACT matching headers
wb_c = openpyxl.Workbook()
ws_c = wb_c.active
ws_c.title = 'Fish Crosses'

header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
header_font = Font(name='Segoe UI', size=11, bold=True, color='38BDF8')

ws_c.append(crosses_field_names)
for cell in ws_c[1]:
    cell.fill = header_fill
    cell.font = header_font

for r in sorted_crosses_rows:
    ws_c.append([clean_val(x) for x in r])

for col in ws_c.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_c.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

wb_c.save(os.path.join(export_dir, 'Crosses_Import_Ready.xlsx'))
wb_c.save(os.path.join(fishnet_dir, 'Crosses_Import_Ready.xlsx'))

with open(os.path.join(export_dir, 'Crosses_Import_Ready.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, lineterminator='\n')
    w.writerow(crosses_field_names)
    for r in sorted_crosses_rows:
        w.writerow([clean_val(x) for x in r])

print(f"Generated Crosses_Import_Ready.xlsx and .csv with {len(sorted_crosses_rows)} rows matching EXACT FileMaker fields!")


# =========================================================================
# 2. TANKS TABLE - EXACT FILEMAKER FIELD NAMES
# =========================================================================
# Exact FileMaker fields:
# TUID, Date of Birth, Date of Death, Turnover Date, Status, Females, Males, Number of Fish, Genotype, Line Nick Name, Notes, Dervitive Cross, Paternal ID, Maternal ID, Father, Mother, Tank Size, Room, Facility, Protocol, Lab Member

tanks_field_names = [
    'TUID',
    'Date of Birth',
    'Date of Death',
    'Turnover Date',
    'Status',
    'Females',
    'Males',
    'Number of Fish',
    'Genotype',
    'Line Nick Name',
    'Notes',
    'Dervitive Cross',
    'Paternal ID',
    'Maternal ID',
    'Father',
    'Mother',
    'Tank Size',
    'Room',
    'Facility',
    'Protocol',
    'Lab Member'
]

tanks_records = {}
for t in fn_tanks:
    tuid = t['TUID'].strip().upper()
    notes = t['NOTES'].strip()
    f_cnt = t['FEMALE'].strip()
    m_cnt = t['MALE'].strip()
    tot = t['TOTAL'].strip()
    dob = t['DOB'].strip()
    dod = t['DOD'].strip()
    turnover = t['TURNOVER'].strip()
    cross = t['Derivative cross'].strip()
    pat = t['PATERNAL'].strip()
    mat = t['MATERNAL'].strip()
    status = t['STATUS'].strip()
    size = t['TANK '].strip()
    
    tanks_records[tuid] = [
        tuid,                                   # TUID
        dob,                                    # Date of Birth
        dod,                                    # Date of Death
        turnover,                               # Turnover Date
        status,                                 # Status
        f_cnt,                                  # Females
        m_cnt,                                  # Males
        tot,                                    # Number of Fish
        get_genotype_desc(notes),               # Genotype
        get_line_nick(notes),                   # Line Nick Name
        notes,                                  # Notes
        cross,                                  # Dervitive Cross
        pat,                                    # Paternal ID
        mat,                                    # Maternal ID
        pat,                                    # Father
        mat,                                    # Mother
        size if size else '1.8L',               # Tank Size
        'D126',                                 # Room
        'Zebrafish facility',                   # Facility
        'QU-IACUC 006/2023-AMM5',               # Protocol
        'Huseyin Cagatay Yalcin'                # Lab Member
    ]

def tuid_sort_key(k):
    m = re.search(r'\d+', k)
    return int(m.group(0)) if m else 9999

sorted_tanks_rows = [tanks_records[k] for k in sorted(tanks_records.keys(), key=tuid_sort_key)]

wb_t = openpyxl.Workbook()
ws_t = wb_t.active
ws_t.title = 'Tanks'

ws_t.append(tanks_field_names)
for cell in ws_t[1]:
    cell.fill = header_fill
    cell.font = header_font

for r in sorted_tanks_rows:
    ws_t.append([clean_val(x) for x in r])

for col in ws_t.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_t.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

wb_t.save(os.path.join(export_dir, 'Tanks_Import_Ready.xlsx'))
wb_t.save(os.path.join(fishnet_dir, 'Tanks_Import_Ready.xlsx'))

with open(os.path.join(export_dir, 'Tanks_Import_Ready.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, lineterminator='\n')
    w.writerow(tanks_field_names)
    for r in sorted_tanks_rows:
        w.writerow([clean_val(x) for x in r])

print(f"Generated Tanks_Import_Ready.xlsx and .csv with {len(sorted_tanks_rows)} rows matching EXACT FileMaker fields!")


# =========================================================================
# 3. NURSERY TABLE - EXACT FILEMAKER FIELD NAMES
# =========================================================================
# Exact FileMaker fields:
# NUID, Derivitive Cross, Number of Fish, Raised, totalFish Graduated, Status, Location

nursery_field_names = [
    'NUID',
    'Derivitive Cross',
    'Number of Fish',
    'Raised',
    'totalFish Graduated',
    'Status',
    'Location'
]

# Read existing nursery
with open(os.path.join(export_dir, 'nursery.tab.orig.bak'), 'r', encoding='utf-8-sig', errors='ignore') as f:
    orig_nursery_raw = list(csv.reader(f, delimiter='\t'))

nursery_records = {}
for r in orig_nursery_raw:
    if len(r) > 0 and r[0].strip():
        nuid = r[0].strip().upper()
        # orig columns: 0:NUID, 1:count, 2:status, 3:tech, 4:CUID, 5:grad_d, 6:entry_d, 7:sp_d, 8:setup_d, 9:mat, 10:pat...
        cuid = r[4] if len(r)>4 else ''
        cnt = r[1] if len(r)>1 else '14'
        stat = r[2] if len(r)>2 else 'Graduate to system'
        nursery_records[nuid] = [
            nuid, cuid, cnt, cnt, cnt, stat, 'D126'
        ]

max_nuid = max([cuid_sort_key(k) for k in nursery_records.keys()] or [0])
existing_n_cuids = set(r[1].upper() for r in nursery_records.values())

for cuid, t_list in cross_to_tanks.items():
    if cuid not in existing_n_cuids:
        max_nuid += 1
        nuid = f"N{max_nuid:04d}"
        tot = str(sum(int(t['TOTAL']) for t in t_list if t['TOTAL'].isdigit()))
        nursery_records[nuid] = [
            nuid, cuid, tot if tot != '0' else '14', tot if tot != '0' else '14', tot if tot != '0' else '14', 'Graduate to system', 'D126'
        ]
        existing_n_cuids.add(cuid)

sorted_nursery_rows = [nursery_records[k] for k in sorted(nursery_records.keys(), key=cuid_sort_key)]

wb_n = openpyxl.Workbook()
ws_n = wb_n.active
ws_n.title = 'Nursery'

ws_n.append(nursery_field_names)
for cell in ws_n[1]:
    cell.fill = header_fill
    cell.font = header_font

for r in sorted_nursery_rows:
    ws_n.append([clean_val(x) for x in r])

for col in ws_n.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_n.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

wb_n.save(os.path.join(export_dir, 'Nursery_Import_Ready.xlsx'))
wb_n.save(os.path.join(fishnet_dir, 'Nursery_Import_Ready.xlsx'))

with open(os.path.join(export_dir, 'Nursery_Import_Ready.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, lineterminator='\n')
    w.writerow(nursery_field_names)
    for r in sorted_nursery_rows:
        w.writerow([clean_val(x) for x in r])

print(f"Generated Nursery_Import_Ready.xlsx and .csv with {len(sorted_nursery_rows)} rows matching EXACT FileMaker fields!")
